"""สร้างโครงการสอนจากแผนการสอน โดยไม่แตะระบบบันทึกหลังสอนใน app.py."""

import io
import json
import copy

import streamlit as st
from online_files import file_uploader as online_file_uploader
from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm
from docxtpl import DocxTemplate
from google import genai
from google.genai import types


FIELDS = {
    "w": "ส.ป.",
    "t": "หัวข้อ",
    "p": "Teaching Point",
    "a": "กิจกรรม",
    "m": "สื่อ",
    "e": "วัดผล",
}


def extract_weeks(api_key, plan_pdf, reference_pdf, count):
    client = genai.Client(api_key=api_key)
    prompt = f"""อ่านแผนการจัดการเรียนรู้ที่แนบมา แล้วจัดทำตารางโครงการสอน {count} สัปดาห์
แผนการสอนเป็นแหล่งข้อมูลหลัก หากมี PDF โครงการสอนตัวอย่าง ให้ใช้อ้างอิงลำดับหัวข้อและรูปแบบตาราง
หากเอกสารสองฉบับขัดกัน ห้ามเดาตัวเลขชั่วโมง ให้ใช้จำนวนสัปดาห์ที่ผู้ใช้ระบุคือ {count}
แต่ละสัปดาห์ต้องมีหัวข้อ Teaching Point กิจกรรม สื่อ และวิธีวัดผลที่สอดคล้องกัน
เขียนช่องวัดผลให้กระชับ ระบุหลักฐานและเกณฑ์สำคัญโดยไม่ใช้ประโยคยาวเกินจำเป็น
ห้ามเขียนผลการเรียนที่เกิดขึ้นแล้ว เพราะเอกสารนี้เป็นโครงการสอนล่วงหน้า
ตอบเป็น JSON เท่านั้น ในรูปแบบ:
{{"weeks":[{{"w":1,"t":"หัวข้อ","p":"สาระสำคัญ","a":"กิจกรรม","m":"สื่อ","e":"วิธีวัดผล"}}]}}
ต้องมีสัปดาห์ 1 ถึง {count} อย่างละหนึ่งรายการ เรียงตามลำดับ
"""
    contents = [types.Part.from_bytes(data=plan_pdf, mime_type="application/pdf")]
    if reference_pdf:
        contents.append(types.Part.from_bytes(data=reference_pdf, mime_type="application/pdf"))
    contents.append(prompt)
    response = client.models.generate_content(
        model="gemini-3.8-flash",
        contents=contents,
        config=types.GenerateContentConfig(response_mime_type="application/json", temperature=0.2),
    )
    if not response.text:
        raise ValueError("AI ไม่ส่งข้อมูลกลับมา")
    data = json.loads(response.text)
    return validate_weeks(data.get("weeks"), count)


def validate_weeks(rows, count):
    if not isinstance(rows, list) or len(rows) != count:
        raise ValueError(f"ตารางต้องมี {count} สัปดาห์พอดี")
    cleaned = []
    for row in rows:
        if not isinstance(row, dict):
            raise ValueError("รูปแบบข้อมูลรายสัปดาห์ไม่ถูกต้อง")
        try:
            week = int(row.get("w"))
        except (ValueError, TypeError):
            raise ValueError("เลขสัปดาห์ต้องเป็นจำนวนเต็ม") from None
        item = {"w": week}
        for key in ("t", "p", "a", "m", "e"):
            item[key] = str(row.get(key) or "").strip()
            if not item[key]:
                raise ValueError(f"สัปดาห์ที่ {week} ยังไม่มีข้อมูลช่อง {FIELDS[key]}")
        cleaned.append(item)
    if sorted(item["w"] for item in cleaned) != list(range(1, count + 1)):
        raise ValueError(f"เลขสัปดาห์ต้องครบ 1–{count} และไม่ซ้ำกัน")
    return sorted(cleaned, key=lambda item: item["w"])


def make_docx(template_bytes, context):
    document = Document(io.BytesIO(template_bytes))
    template_text = "\n".join(
        cell.text
        for table in document.tables
        for row in table.rows
        for cell in row.cells
    )
    if "{%tr for r in weeks %}" not in template_text or "{%tr endfor %}" not in template_text:
        raise ValueError("แบบฟอร์มต้องมีแถว {%tr for r in weeks %} และ {%tr endfor %}")
    if "{{ ... }}" in template_text:
        raise ValueError("พบ {{ ... }} ในแบบฟอร์ม กรุณาเปลี่ยนเป็นเลขหน้า PAGE ของ Word")
    template = DocxTemplate(io.BytesIO(template_bytes))
    template.render(context)
    result = io.BytesIO()
    template.save(result)
    document = Document(io.BytesIO(result.getvalue()))
    if document.tables:
        table = document.tables[0]
        while len(table.rows) > 6 and all(not cell.text.strip() for cell in table.rows[-1].cells):
            table._tbl.remove(table.rows[-1]._tr)
        if len(table._tbl.tblGrid.gridCol_lst) == 9:
            widths = [Cm(value).twips for value in (1.2, 1.0, 1.9, 3.1, 3.55, 0.05, 2.25, 0.15, 3.8)]
            for column, width in zip(table._tbl.tblGrid.gridCol_lst, widths):
                column.w = width * 635
            for row in table._tbl.tr_lst:
                offset = 0
                for cell in row.tc_lst:
                    span = cell.grid_span
                    cell.tcPr.get_or_add_tcW().w = sum(widths[offset:offset + span])
                    cell.tcPr.tcW.type = "dxa"
                    offset += span
        seen = set()
        for row in table.rows[:6]:
            for cell in row.cells:
                if cell._tc in seen:
                    continue
                seen.add(cell._tc)
                if not ("หน้าที่" in cell.text or "แผ่นที่" in cell.text):
                    continue
                if cell._tc.xpath('.//w:fldSimple') or cell._tc.xpath('.//w:instrText'):
                    continue
                paragraph = cell.paragraphs[-1] if "หน้าที่" in cell.text else cell.paragraphs[0]
                field = OxmlElement("w:fldSimple")
                field.set(qn("w:instr"), "PAGE")
                run = OxmlElement("w:r")
                if paragraph.runs and paragraph.runs[-1]._r.rPr is not None:
                    run.append(copy.deepcopy(paragraph.runs[-1]._r.rPr))
                value = OxmlElement("w:t")
                value.text = "1"
                run.append(value)
                field.append(run)
                paragraph._p.append(field)
        settings = document.settings._element
        if not settings.xpath("./w:updateFields"):
            update = OxmlElement("w:updateFields")
            update.set(qn("w:val"), "true")
            settings.append(update)
        # A real Word page header repeats automatically while the body table
        # continues until the page is full. PAGE fields then update per page.
        heading = copy.deepcopy(table._tbl)
        for row in list(heading.tr_lst)[6:]:
            heading.remove(row)
        for row in list(table._tbl.tr_lst)[:6]:
            table._tbl.remove(row)
        section = document.sections[0]
        section.header_distance = Cm(2.5)
        section.top_margin = Cm(6.1)
        header_part = section.header.part
        for image in heading.xpath('.//a:blip[@r:embed]'):
            source_id = image.get(qn("r:embed"))
            image.set(qn("r:embed"), header_part.relate_to(document.part.related_parts[source_id], RT.IMAGE))
        section.header._element.insert(0, heading)
        for paragraph in section.header._element.xpath("./w:p"):
            if not "".join(paragraph.itertext()).strip():
                section.header._element.remove(paragraph)
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


st.set_page_config(
    page_title="ระบบจัดทำโครงการสอน AI",
    page_icon="📘",
    layout="wide",
    initial_sidebar_state="expanded",
)
st.html("""
<style>
@import url('https://fonts.googleapis.com/css2?family=Prompt:wght@300;400;500;600;700&display=swap');
html, body, [data-testid="stAppViewContainer"] { font-family: 'Prompt', sans-serif; }
[data-testid="stAppViewContainer"] { background: #f6f9ff; }
.project-hero {
    background: linear-gradient(125deg, #183b88, #336be4 58%, #08b9cc);
    border-radius: 20px; padding: 32px 38px; color: white;
    box-shadow: 0 14px 32px rgba(30, 58, 138, .18);
    margin-bottom: 22px;
}
.project-hero .eyebrow { font-size: 12px; letter-spacing: .13em; font-weight: 600; opacity: .85; }
.project-hero h1 { color: white; margin: 10px 0 8px; font-size: clamp(28px, 3vw, 42px); line-height: 1.25; }
.project-hero p { color: #e6f5ff; font-size: 15px; margin: 0; }
.project-steps { display: flex; flex-wrap: wrap; gap: 9px; margin: 0 0 18px; }
.project-steps span { background: #e9f1ff; color: #214484; border-radius: 999px; padding: 7px 14px; font-size: 13px; font-weight: 500; }
[data-testid="stVerticalBlockBorderWrapper"] { border-radius: 16px; background: white; }
.stButton > button[kind="primary"], .stDownloadButton > button {
    background: linear-gradient(90deg, #f43f72, #df235b); color: white;
    border: 0; border-radius: 12px; font-weight: 600;
    box-shadow: 0 7px 18px rgba(225, 40, 90, .16);
}
.stButton > button[kind="primary"]:hover, .stDownloadButton > button:hover {
    color: white; border: 0; background: #c91d50;
}
</style>
""")
if st.session_state.get("project_output_version") != 4:
    st.session_state.pop("project_output", None)
st.html("""
<div class="project-hero">
  <div class="eyebrow">AI VOCATIONAL PLANNING</div>
  <h1>ระบบจัดทำโครงการสอน</h1>
  <p>เปลี่ยนแผนการสอนเป็นตารางรายสัปดาห์ ตรวจแก้ข้อมูล และส่งออกตามแบบฟอร์มวิทยาลัย</p>
</div>
<div class="project-steps">
  <span>01 · แนบเอกสาร</span><span>02 · กำหนดข้อมูลวิชา</span>
  <span>03 · ตรวจแก้รายสัปดาห์</span><span>04 · ดาวน์โหลด Word</span>
</div>
""")

with st.sidebar:
    if st.session_state.get("combined_app"):
        st.page_link("app.py", label="ไปยังระบบบันทึกหลังสอน", icon=":material/swap_horiz:", width="stretch")
    else:
        st.link_button(
            "ไปยังระบบบันทึกหลังสอน",
            "https://smart-teaching-log-system-3bcqd8atj7upyy3twdx7ed.streamlit.app/",
            icon=":material/swap_horiz:",
            width="stretch",
        )
    st.divider()
    st.header("ตั้งค่าระบบ", icon=":material/settings:")
    st.subheader("เชื่อมต่อ Gemini")
    api_key = st.text_input("Gemini API Key", type="password", placeholder="กรอก API Key เพื่อวิเคราะห์ PDF")
    st.caption("เอกสารจะถูกส่งให้ Gemini เมื่อกดวิเคราะห์เท่านั้น")
    st.markdown("[รับ API Key จาก Google AI Studio](https://aistudio.google.com/apikey)")
    st.divider()
    st.caption("ระบบจัดทำโครงการสอน · ส่งออก Word ตามแบบฟอร์มวิทยาลัย")

left, right = st.columns(2, gap="large")
with left:
    with st.container(border=True):
        st.subheader("1. เอกสารอ้างอิง", icon=":material/upload_file:")
        st.caption("แผนการสอนเป็นข้อมูลหลัก ส่วนโครงการสอนตัวอย่างช่วยกำหนดลำดับและรูปแบบ")
        if st.session_state.get("combined_app"):
            plan_file = online_file_uploader("แผนการสอน PDF", ["pdf"], "project_plan")
            reference_file = online_file_uploader("โครงการสอนตัวอย่าง PDF (ถ้ามี)", ["pdf"], "project_reference")
            template_file = online_file_uploader("แบบฟอร์มโครงการสอน .docx", ["docx"], "project_template")
        else:
            plan_file = st.file_uploader("แผนการสอน PDF", type="pdf")
            reference_file = st.file_uploader("โครงการสอนตัวอย่าง PDF (ถ้ามี)", type="pdf")
            template_file = st.file_uploader("แบบฟอร์มโครงการสอน .docx", type="docx")
with right:
    with st.container(border=True):
        st.subheader("2. ข้อมูลรายวิชา", icon=":material/school:")
        code_col, degree_col, year_col = st.columns([1.3, 1, 1])
        with code_col:
            code = st.text_input("รหัสวิชา", value="21901-2011")
        with degree_col:
            degree = st.selectbox("ระดับ", ["ปวช.", "ปวส."])
        with year_col:
            year_level = st.selectbox("ปีที่", [1, 2, 3] if degree == "ปวช." else [1, 2])
        subject = st.text_input("ชื่อวิชา", value="การพัฒนาแอปพลิเคชันบนอุปกรณ์เคลื่อนที่")
        hours_col, weeks_col, term_col = st.columns(3)
        with hours_col:
            h = st.number_input("ชั่วโมง/สัปดาห์", min_value=1, value=5)
        with weeks_col:
            n = st.number_input("สัปดาห์/ภาคเรียน", min_value=1, max_value=40, value=18 if degree == "ปวช." else 15)
        with term_col:
            term = st.text_input("ภาคเรียนที่", value="1/2569")

curriculum = (
    "ประกาศนียบัตรวิชาชีพ พ.ศ. 2567"
    if degree == "ปวช."
    else "ประกาศนียบัตรวิชาชีพชั้นสูง พ.ศ. 2567"
)

if st.button("วิเคราะห์แผนการสอน", type="primary", icon=":material/auto_awesome:", width="stretch"):
    if not api_key or not plan_file:
        st.error("กรุณากรอก Gemini API Key และแนบแผนการสอน PDF")
    else:
        try:
            with st.spinner("กำลังอ่านแผนการสอนและจัดตารางรายสัปดาห์..."):
                st.session_state.project_rows = extract_weeks(
                    api_key,
                    plan_file.getvalue(),
                    reference_file.getvalue() if reference_file else None,
                    int(n),
                )
                st.session_state.pop("project_output", None)
            st.success("อ่านข้อมูลแล้ว กรุณาตรวจแก้ทุกแถวก่อนสร้างเอกสาร")
        except Exception as exc:
            error_msg = str(exc)
            if "403" in error_msg and "PERMISSION_DENIED" in error_msg:
                st.error("⚠️ วิเคราะห์ไม่สำเร็จ: API Key ของคุณมีปัญหา หรือถูกระงับการใช้งาน โปรดสร้าง API Key ใหม่ที่ Google AI Studio และนำมาอัปเดตใหม่ครับ")
            else:
                st.error(f"วิเคราะห์ไม่สำเร็จ: {exc}")

if "project_rows" in st.session_state:
    with st.container(border=True):
        st.subheader("3. ตรวจแก้โครงการสอน", icon=":material/edit_document:")
        st.caption("ตรวจความถูกต้องของแต่ละสัปดาห์ แล้วกดสร้างเอกสาร Word")
        edited = st.data_editor(
            st.session_state.project_rows,
            column_config={
                "w": st.column_config.NumberColumn("ส.ป.", min_value=1, step=1),
                **{key: st.column_config.TextColumn(label) for key, label in FIELDS.items() if key != "w"},
            },
            num_rows="dynamic",
            hide_index=True,
            width="stretch",
        )
        create_col, download_col = st.columns([1, 1])
        with create_col:
            create_document = st.button("สร้างไฟล์ Word", icon=":material/description:", width="stretch")
        if create_document:
            if not template_file:
                st.error("กรุณาแนบแบบฟอร์มโครงการสอน .docx")
            else:
                try:
                    rows = edited.to_dict("records") if hasattr(edited, "to_dict") else edited
                    rows = validate_weeks(rows, int(n))
                    context = {
                        "curriculum": curriculum,
                        "code": code.strip(),
                        "subject": subject.strip(),
                        "level": "ประกาศนียบัตรวิชาชีพ" if degree == "ปวช." else "ประกาศนียบัตรวิชาชีพชั้นสูง",
                        "year_level": year_level,
                        "h": h,
                        "n": n,
                        "term": term.strip(),
                        "weeks": rows,
                    }
                    output = make_docx(template_file.getvalue(), context)
                    st.session_state.project_output = output
                    st.session_state.project_output_version = 4
                    st.session_state.project_filename = f"โครงการสอน_{code.strip()}.docx"
                    st.success("เอกสารพร้อมดาวน์โหลด")
                except Exception as exc:
                    st.error(f"สร้างเอกสารไม่สำเร็จ: {exc}")
        with download_col:
            if "project_output" in st.session_state:
                st.download_button(
                    "ดาวน์โหลดโครงการสอน .docx",
                    data=st.session_state.project_output,
                    file_name=st.session_state.project_filename,
                    mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                    icon=":material/download:",
                    width="stretch",
                )
